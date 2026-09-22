from transformers import AutoTokenizer, AutoModel
import openpyxl
import os
import numpy as np
import torch
import argparse


# ============================================================
# 1. Load tokenizer and BERT model
# ============================================================

tokenizer = AutoTokenizer.from_pretrained(
    "bert-base-chinese"
)

model = AutoModel.from_pretrained(
    "bert-base-chinese",
    output_hidden_states=True
)

model.eval()


# ============================================================
# 2. Novel configuration (cross-platform paths)
# ============================================================

# 数据集根目录：
#   1. 优先读取环境变量 CHINESE_EEG_ROOT（Windows/Linux 通用）
#   2. Windows 默认： D:/AI/ChineseEEG
#   3. Linux 默认：   ~/ChineseEEG
if "CHINESE_EEG_ROOT" in os.environ:
    EEG_ROOT = os.environ["CHINESE_EEG_ROOT"]
elif os.name == "nt":
    EEG_ROOT = "D:/AI/ChineseEEG"
else:
    EEG_ROOT = os.path.expanduser("~/ChineseEEG")

# 输出根目录：固定为脚本所在目录的 ../outputs，
# 不依赖运行时的工作目录（cwd）
OUTPUT_ROOT = os.path.normpath(
    os.path.join(
        os.path.dirname(os.path.abspath(__file__)),
        "..",
        "outputs"
    )
)

Chinese_novels = {
    "LittlePrince": {
        "segmented_path": os.path.join(
            EEG_ROOT,
            "derivatives",
            "novels",
            "segmented_novel",
            "LittlePrince"
        ),
        "run_num": 7,
        "embedding_path": os.path.join(
            OUTPUT_ROOT,
            "embeddings",
            "LittlePrince"
        ),
    },

    "GarnettDream": {
        "segmented_path": os.path.join(
            EEG_ROOT,
            "derivatives",
            "novels",
            "segmented_novel",
            "GarnettDream"
        ),
        "run_num": 18,
        "embedding_path": os.path.join(
            OUTPUT_ROOT,
            "embeddings",
            "GarnettDream"
        ),
    }
}


# ============================================================
# 3. Parse arguments
# ============================================================

parser = argparse.ArgumentParser(
    description="Extract BERT sentence embeddings"
)

parser.add_argument(
    "--novel_name",
    type=str,
    default="LittlePrince",
    help="Novel key in Chinese_novels (e.g. LittlePrince, GarnettDream)"
)

args = parser.parse_args()


# ============================================================
# 4. Load novel configuration
# ============================================================

novel_cfg = Chinese_novels[args.novel_name]

args.segmented_path = novel_cfg["segmented_path"]
args.run_num = novel_cfg["run_num"]
args.embedding_path = novel_cfg["embedding_path"]

os.makedirs(
    args.embedding_path,
    exist_ok=True
)


# ============================================================
# 5. Extract embeddings for each run
# ============================================================

for i in range(args.run_num):

    novel_path = os.path.join(
        args.segmented_path,
        # 注意：数据文件名本身拼写为 "Chinense"（非 "Chinese"）
        f"segmented_Chinense_novel_run_{i + 1}.xlsx"
    )

    # --------------------------------------------------------
    # Read Excel
    # --------------------------------------------------------

    wb = openpyxl.load_workbook(novel_path)
    wsheet = wb.active

    texts = []

    for j in range(2, wsheet.max_row + 1):

        text = wsheet.cell(
            row=j,
            column=1
        ).value

        # Skip empty cells
        if text is not None:
            texts.append(str(text))


    # --------------------------------------------------------
    # Extract sentence embeddings
    # --------------------------------------------------------

    embeddings = []

    for k, text in enumerate(texts):

        # text
        #   ↓
        # [CLS] token1 token2 ... tokenN [SEP]
        #   ↓
        # input_ids

        inputs = tokenizer(
            text,
            return_tensors="pt",
            return_special_tokens_mask=True
        )

        with torch.no_grad():

            outputs = model(
                input_ids=inputs["input_ids"],
                attention_mask=inputs["attention_mask"]
            )

            # outputs.hidden_states:
            #
            # length = 13
            #
            # hidden_states[0]
            #     embedding layer
            #
            # hidden_states[1]
            #     Transformer Layer 1
            #
            # ...
            #
            # hidden_states[12]
            #     Transformer Layer 12


            # ------------------------------------------------
            # Last 4 Transformer layers
            # ------------------------------------------------

            last_four = torch.stack(
                outputs.hidden_states[-4:],
                dim=0
            )

            # shape:
            #
            # (4, batch_size, seq_len, 768)


            # ------------------------------------------------
            # Average last 4 layers
            # ------------------------------------------------

            hidden = last_four.mean(dim=0)

            # shape:
            #
            # (batch_size, seq_len, 768)


            # ------------------------------------------------
            # Remove special tokens
            #
            # [CLS] ×
            # normal token ✓
            # [SEP] ×
            # [PAD] ×
            # ------------------------------------------------

            mask = (
                inputs["attention_mask"]
                * (1 - inputs["special_tokens_mask"])
            )

            # (batch_size, seq_len)

            mask = mask.unsqueeze(-1)

            # (batch_size, seq_len, 1)


            # ------------------------------------------------
            # Mean pooling over normal text tokens
            # ------------------------------------------------

            sentence_embedding = (
                (hidden * mask).sum(dim=1)
                / mask.sum(dim=1)
            )

            # shape:
            #
            # (1, 768)


        # Remove batch dimension:
        #
        # (1, 768)
        #     ↓
        # (768,)

        sentence_embedding = (
            sentence_embedding
            .squeeze(0)
            .cpu()
            .numpy()
        )

        embeddings.append(
            sentence_embedding
        )


    # --------------------------------------------------------
    # List -> numpy array
    # --------------------------------------------------------

    embeddings = np.stack(
        embeddings,
        axis=0
    )

    # shape:
    #
    # (num_texts, 768)


    # --------------------------------------------------------
    # Save
    # --------------------------------------------------------

    save_path = os.path.join(
        args.embedding_path,
        f"text_embedding_run_{i + 1}.npy"
    )

    np.save(
        save_path,
        embeddings
    )

    print(
        f"Run {i + 1}: "
        f"{len(texts)} texts -> "
        f"{embeddings.shape}, "
        f"saved to {save_path}"
    )