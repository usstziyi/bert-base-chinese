from transformers import AutoTokenizer, AutoModel
import openpyxl
import os
import numpy as np
import torch
import argparse

# 数据集根目录：
EEG_ROOT = "data/ChineseEEG"
# 原始小说目录
ORI_NOVEL_PATH = os.path.join(EEG_ROOT, "derivatives", "novels", "original_novel")
# 分段小说目录
SEG_NOVEL_PATH = os.path.join(EEG_ROOT, "derivatives", "novels", "segmented_novel")



def load_text(novel_name='LittlePrince', run_num=7):
    segmented_path = os.path.join(SEG_NOVEL_PATH, novel_name)

    text_data = []
    # 遍历每个run
    for i in range(run_num):
        novel_path = os.path.join(
            segmented_path,
            f"segmented_Chinense_novel_run_{i + 1}.xlsx"
        )


        # Read Excel file
        wb = openpyxl.load_workbook(novel_path, read_only=True, data_only=True)
        wsheet = wb.active

        run_texts = []
        for j in range(2, wsheet.max_row + 1):
            text = wsheet.cell(row=j, column=1).value
            # Skip: None / "" / "   "
            if text is not None and str(text).strip():
                run_texts.append(str(text).strip())

        wb.close()

        text_data.append(run_texts)

    return text_data


if __name__ == "__main__":
    text_data = load_text(novel_name='LittlePrince', run_num=7)
    print(len(text_data))
    for run_texts in text_data:
        print(len(run_texts))
